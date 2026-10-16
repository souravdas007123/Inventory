"""One Excel file for the accountant (CA): sales and purchase registers, credit / debit notes, expenses, payments.
Every sheet is read in a few grouped database queries (nothing is calculated per bill), so a whole year is fast."""
import io
from collections import OrderedDict
from decimal import Decimal

from django.db.models import Count, Sum

from .models import (
    Bill, CustomerPayment, Expense, InvoiceItem, Payment, Purchase, PurchaseItem, PurchaseReturn,
    PurchaseReturnItem, SalesReturn, SalesReturnItem,
)

ZERO = Decimal('0')
MAX_DAYS = 800                    # a range longer than about 2 years is refused (keeps the file and the wait small)
MAX_ROWS = 100000                 # most rows in one sheet (a bigger range is refused, never cut off)


def d2(x):
    return round(float(x or 0), 2)


def dstr(x):
    return x.strftime('%d-%m-%Y') if x else ''


class RangeError(ValueError):
    pass


def check_range(start, end):
    if start > end:
        raise RangeError('The "From" date is after the "To" date.')
    if (end - start).days > MAX_DAYS:
        raise RangeError('Please choose a range of up to 2 years.')
    check_size(start, end)


def check_size(start, end):
    """Refuses a range with more rows than a sheet should hold (so nothing is ever cut off silently)."""
    counts = (Bill.objects.filter(date__gte=start, date__lte=end).count(),
              Purchase.objects.filter(invoice_date__gte=start, invoice_date__lte=end).count(),
              Expense.objects.filter(date__gte=start, date__lte=end).count(),
              Payment.objects.filter(payment_date__gte=start, payment_date__lte=end).count(),
              CustomerPayment.objects.filter(date__gte=start, date__lte=end).count())
    if max(counts) > MAX_ROWS:
        raise RangeError('There are too many entries in this range. Please choose a shorter range.')


def _sums(qs, key, **fields):
    """{key value: {name: Decimal}} with one grouped query."""
    out = OrderedDict()
    for r in qs.values(key).annotate(**{n: Sum(f) for n, f in fields.items()}).order_by():
        out[r[key]] = {n: r[n] or ZERO for n in fields}
    return out


def sales_data(start, end):
    bills = list(Bill.objects.filter(status='FINAL', date__gte=start, date__lte=end)
                 .select_related('customer').order_by('date', 'id'))
    sums = _sums(InvoiceItem.objects.filter(bill__status='FINAL', bill__date__gte=start, bill__date__lte=end), 'bill_id',
                 taxable='taxable_value', cgst='cgst', sgst='sgst', igst='igst', total='total')
    rows = []
    for b in bills:
        s = sums.get(b.pk) or dict(taxable=ZERO, cgst=ZERO, sgst=ZERO, igst=ZERO, total=ZERO)
        rows.append(dict(date=b.date, invoice=b.display_number, customer=b.customer_name or '',
                         gstin=(b.customer.gstin if b.customer_id and b.customer.gstin else ''),
                         payment=b.get_payment_mode_display(), **s))
    by_rate = _sums(InvoiceItem.objects.filter(bill__status='FINAL', bill__date__gte=start, bill__date__lte=end), 'gst_rate',
                    taxable='taxable_value', cgst='cgst', sgst='sgst', igst='igst', total='total')
    cancelled = Bill.objects.filter(status='CANCELLED', date__gte=start, date__lte=end).count()
    return rows, by_rate, cancelled


def credit_notes(start, end):
    qs = SalesReturnItem.objects.filter(sales_return__date__gte=start, sales_return__date__lte=end)
    sums = _sums(qs, 'sales_return_id', taxable='taxable_value', cgst='cgst', sgst='sgst', igst='igst', total='total')
    rows = []
    for r in SalesReturn.objects.filter(pk__in=sums).select_related('bill', 'bill__customer').order_by('date', 'id'):
        c = r.bill.customer
        rows.append(dict(date=r.date, note=r.note_number or f'CN-{r.pk}', invoice=r.bill.display_number, customer=r.bill.customer_name or '',
                         gstin=(c.gstin if c and c.gstin else ''), reason=r.reason, **sums[r.pk]))
    return rows


def purchase_data(start, end):
    purchases = list(Purchase.objects.filter(invoice_date__gte=start, invoice_date__lte=end)
                     .select_related('supplier').order_by('invoice_date', 'id'))
    base = PurchaseItem.objects.filter(purchase__invoice_date__gte=start, purchase__invoice_date__lte=end)
    sums = _sums(base, 'purchase_id', taxable='tax', gst='gst', cgst='cgst', sgst='sgst', total='purchase_price')
    rows = []
    for p in purchases:
        s = sums.get(p.pk) or dict(taxable=ZERO, gst=ZERO, cgst=ZERO, sgst=ZERO, total=ZERO)
        rows.append(dict(date=p.invoice_date, invoice=p.invoice_number, supplier=p.supplier.company,
                         gstin=getattr(p.supplier, 'gstin', '') or '', **s))
    by_rate = _sums(base, 'gst_rate', taxable='tax', gst='gst', cgst='cgst', sgst='sgst', total='purchase_price')
    return rows, by_rate


def debit_notes(start, end):
    qs = PurchaseReturnItem.objects.filter(purchase_return__date__gte=start, purchase_return__date__lte=end)
    sums = _sums(qs, 'purchase_return_id', taxable='taxable_value', gst='gst', total='amount')
    rows = []
    for r in PurchaseReturn.objects.filter(pk__in=sums).select_related('purchase', 'purchase__supplier').order_by('date', 'id'):
        sup = r.purchase.supplier
        rows.append(dict(date=r.date, note=r.note_number or f'DN-{r.pk}', invoice=r.purchase.invoice_number, supplier=sup.company,
                         gstin=getattr(sup, 'gstin', '') or '', reason=r.reason, **sums[r.pk]))
    return rows


def totals(rows, *fields):
    return {f: sum((r[f] for r in rows), ZERO) for f in fields}


def summary(start, end):
    """The numbers of the Summary sheet (also used by the page and the tests)."""
    s_rows, _, cancelled = sales_data(start, end)
    cn = credit_notes(start, end)
    p_rows, _ = purchase_data(start, end)
    dn = debit_notes(start, end)
    S = totals(s_rows, 'taxable', 'cgst', 'sgst', 'igst', 'total')
    C = totals(cn, 'taxable', 'cgst', 'sgst', 'igst', 'total')
    P = totals(p_rows, 'taxable', 'gst', 'cgst', 'sgst', 'total')
    D = totals(dn, 'taxable', 'gst', 'total')
    exp = Expense.objects.filter(date__gte=start, date__lte=end).aggregate(s=Sum('amount'), n=Count('id'))
    pay = Payment.objects.filter(payment_date__gte=start, payment_date__lte=end).aggregate(s=Sum('amount'))
    rec = CustomerPayment.objects.filter(date__gte=start, date__lte=end).aggregate(s=Sum('amount'))
    S['gst_total'] = S['cgst'] + S['sgst'] + S['igst']
    C['gst_total'] = C['cgst'] + C['sgst'] + C['igst']
    return dict(sales=S, credit=C, purchases=P, debit=D, bills=len(s_rows), cancelled=cancelled, credit_notes=len(cn),
                purchase_count=len(p_rows), debit_notes=len(dn), expenses=exp['s'] or ZERO, supplier_paid=pay['s'] or ZERO,
                received=rec['s'] or ZERO)


def build_workbook(start, end):
    """The Excel file (bytes) for start..end (both days included)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    check_range(start, end)
    s_rows, by_rate, cancelled = sales_data(start, end)
    cn = credit_notes(start, end)
    p_rows, p_rate = purchase_data(start, end)
    dn = debit_notes(start, end)
    S = totals(s_rows, 'taxable', 'cgst', 'sgst', 'igst', 'total')
    C = totals(cn, 'taxable', 'cgst', 'sgst', 'igst', 'total')
    P = totals(p_rows, 'taxable', 'gst', 'cgst', 'sgst', 'total')
    D = totals(dn, 'taxable', 'gst', 'total')
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, headers, rows, money_from, total_label=None, total_cols=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)
        if total_label is not None and rows:
            n = len(rows) + 1
            line = [''] * len(headers)
            line[0] = total_label
            for i in total_cols:
                col = ws.cell(row=1, column=i).column_letter
                line[i - 1] = f'=SUM({col}2:{col}{n})'
            ws.append(line)
            for c in ws[ws.max_row]:
                c.font = bold
        for i in range(money_from, len(headers) + 1):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = '#,##0.00'
        for i, h in enumerate(headers, 1):
            w = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if len(r) >= i] or [10])
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(w + 2, 11), 42)
        ws.freeze_panes = 'A2'
        return ws

    net_sales = {k: S[k] - C[k] for k in ('taxable', 'cgst', 'sgst', 'igst', 'total')}
    net_p_gst, net_p_taxable = P['gst'] - D['gst'], P['taxable'] - D['taxable']
    out_gst = net_sales['cgst'] + net_sales['sgst'] + net_sales['igst']
    exp = Expense.objects.filter(date__gte=start, date__lte=end).aggregate(s=Sum('amount'))['s'] or ZERO
    paid = Payment.objects.filter(payment_date__gte=start, payment_date__lte=end).aggregate(s=Sum('amount'))['s'] or ZERO
    got = CustomerPayment.objects.filter(date__gte=start, date__lte=end).aggregate(s=Sum('amount'))['s'] or ZERO
    sheet('Summary', ['Particular', 'Taxable', 'CGST', 'SGST', 'IGST', 'Total'], [
        [f'Period: {dstr(start)} to {dstr(end)}', '', '', '', '', ''],
        ['Sales (final bills)', d2(S['taxable']), d2(S['cgst']), d2(S['sgst']), d2(S['igst']), d2(S['total'])],
        ['less: Credit notes (sales returns)', d2(C['taxable']), d2(C['cgst']), d2(C['sgst']), d2(C['igst']), d2(C['total'])],
        ['Net sales', d2(net_sales['taxable']), d2(net_sales['cgst']), d2(net_sales['sgst']), d2(net_sales['igst']), d2(net_sales['total'])],
        ['Purchases', d2(P['taxable']), d2(P['cgst']), d2(P['sgst']), '', d2(P['total'])],
        ['less: Debit notes (purchase returns)', d2(D['taxable']), d2(D['gst'] / 2), d2(D['gst'] - D['gst'] / 2), '', d2(D['total'])],
        ['Net purchases (GST = input credit, CGST + SGST split equally)', d2(net_p_taxable), d2(net_p_gst / 2), d2(net_p_gst - net_p_gst / 2), '', d2(P['total'] - D['total'])],
        ['GST payable = output GST - input GST', '', '', '', '', d2(out_gst - net_p_gst)],
        ['', '', '', '', '', ''],
        ['Number of final bills', '', '', '', '', len(s_rows)],
        ['Number of cancelled bills (not counted above)', '', '', '', '', cancelled],
        ['Number of credit notes', '', '', '', '', len(cn)],
        ['Number of purchase invoices', '', '', '', '', len(p_rows)],
        ['Number of debit notes', '', '', '', '', len(dn)],
        ['Expenses', '', '', '', '', d2(exp)],
        ['Paid to suppliers', '', '', '', '', d2(paid)],
        ['Received from credit customers', '', '', '', '', d2(got)],
        ['', '', '', '', '', ''],
        ['This is an accounting summary. Please verify it with your CA before filing a return.', '', '', '', '', ''],
    ], 2, first=True)
    sheet('Sales Register', ['Date', 'Invoice', 'Customer', 'GSTIN', 'Payment', 'Taxable', 'CGST', 'SGST', 'IGST', 'Total'],
          [[dstr(r['date']), r['invoice'], r['customer'], r['gstin'], r['payment'], d2(r['taxable']), d2(r['cgst']), d2(r['sgst']),
            d2(r['igst']), d2(r['total'])] for r in s_rows], 6, 'Total', range(6, 11))
    sheet('Sales by GST rate', ['GST %', 'Taxable', 'CGST', 'SGST', 'IGST', 'Total'],
          [[k, d2(v['taxable']), d2(v['cgst']), d2(v['sgst']), d2(v['igst']), d2(v['total'])] for k, v in sorted(by_rate.items())], 2)
    sheet('Credit Notes', ['Date', 'Note', 'Against invoice', 'Customer', 'GSTIN', 'Reason', 'Taxable', 'CGST', 'SGST', 'IGST', 'Total'],
          [[dstr(r['date']), r['note'], r['invoice'], r['customer'], r['gstin'], r['reason'], d2(r['taxable']), d2(r['cgst']), d2(r['sgst']),
            d2(r['igst']), d2(r['total'])] for r in cn], 7, 'Total', range(7, 12))
    sheet('Purchase Register', ['Date', 'Invoice', 'Supplier', 'GSTIN', 'Taxable', 'CGST', 'SGST', 'GST', 'Total'],
          [[dstr(r['date']), r['invoice'], r['supplier'], r['gstin'], d2(r['taxable']), d2(r['cgst']), d2(r['sgst']), d2(r['gst']),
            d2(r['total'])] for r in p_rows], 5, 'Total', range(5, 10))
    sheet('Purchases by GST rate', ['GST %', 'Taxable', 'CGST', 'SGST', 'GST', 'Total'],
          [[k, d2(v['taxable']), d2(v['cgst']), d2(v['sgst']), d2(v['gst']), d2(v['total'])] for k, v in sorted(p_rate.items())], 2)
    sheet('Debit Notes', ['Date', 'Note', 'Against invoice', 'Supplier', 'GSTIN', 'Reason', 'Taxable', 'GST', 'Total'],
          [[dstr(r['date']), r['note'], r['invoice'], r['supplier'], r['gstin'], r['reason'], d2(r['taxable']), d2(r['gst']), d2(r['total'])]
           for r in dn], 7, 'Total', range(7, 10))
    ex = [[dstr(e.date), e.get_name_display(), e.employee.name if e.employee_id else '', e.note, d2(e.amount)]
          for e in Expense.objects.filter(date__gte=start, date__lte=end).select_related('employee').order_by('date', 'id')]
    sheet('Expenses', ['Date', 'Type', 'Employee', 'Note', 'Amount'], ex, 5, 'Total', [5])
    sp = [[dstr(p.payment_date), p.supplier.company, p.get_payment_mode_display() if p.payment_mode else '', d2(p.amount)]
          for p in Payment.objects.filter(payment_date__gte=start, payment_date__lte=end).select_related('supplier').order_by('payment_date', 'id')]
    sheet('Supplier Payments', ['Date', 'Supplier', 'Mode', 'Amount'], sp, 4, 'Total', [4])
    cr = [[dstr(p.date), p.customer.name, p.get_mode_display(), p.note, d2(p.amount)]
          for p in CustomerPayment.objects.filter(date__gte=start, date__lte=end).select_related('customer').order_by('date', 'id')]
    sheet('Customer Receipts', ['Date', 'Customer', 'Mode', 'Note', 'Amount'], cr, 5, 'Total', [5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ['build_workbook', 'check_range', 'RangeError', 'summary']
