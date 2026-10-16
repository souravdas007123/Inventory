"""The Excel file for the accountant."""
import io
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from backend import ca_export
from backend.models import Customer, Expense
from backend.models_extra import PurchaseReturn, PurchaseReturnItem, SalesReturn, SalesReturnItem

from .helpers import StockMixin, today


def rows(wb, name):
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


@override_settings(SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False, CSRF_COOKIE_SECURE=False)
class CaExportTests(StockMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.owner = get_user_model().objects.create_superuser('owner', 'o@example.com', 'pw-12345')
        self.client.force_login(self.owner)
        self.purchase, self.pitem = self.buy(qty=5, price='560.00')          # 5 Strip, taxable 500, GST 60
        self.bill, self.item = self.sell(qty=10, rate='10.00')               # taxable 100, GST 12, total 112
        self.sr = SalesReturn.objects.create(bill=self.bill, date=today())
        SalesReturnItem.objects.create(sales_return=self.sr, invoice_item=self.item, qty=4)     # taxable 40, GST 4.80, total 44.80
        self.start, self.end = today() - timedelta(days=5), today()

    def book(self, start=None, end=None):
        return load_workbook(io.BytesIO(ca_export.build_workbook(start or self.start, end or self.end)))

    def test_sheets_and_the_sales_register(self):
        wb = self.book()
        self.assertEqual(wb.sheetnames, ['Summary', 'Sales Register', 'Sales by GST rate', 'Credit Notes', 'Purchase Register',
                                         'Purchases by GST rate', 'Debit Notes', 'Expenses', 'Supplier Payments', 'Customer Receipts'])
        sales = rows(wb, 'Sales Register')
        self.assertEqual(sales[0][:3], ['Date', 'Invoice', 'Customer'])
        self.assertEqual(sales[1][1], self.bill.display_number)
        self.assertEqual(sales[1][5:], [100.0, 6.0, 6.0, 0.0, 112.0])
        self.assertTrue(str(sales[2][5]).startswith('=SUM('))                # a total line

    def test_credit_note_and_gst_rate_sheets(self):
        wb = self.book()
        cn = rows(wb, 'Credit Notes')
        self.assertTrue(cn[1][1].startswith('CN/'))
        self.assertEqual(cn[1][6:], [40.0, 2.4, 2.4, 0.0, 44.8])
        self.assertEqual(rows(wb, 'Sales by GST rate')[1], [12, 100.0, 6.0, 6.0, 0.0, 112.0])

    def test_purchase_register(self):
        wb = self.book()
        pur = rows(wb, 'Purchase Register')
        self.assertEqual(pur[1][1:3], [self.purchase.invoice_number, 'ABC Pharma'])
        self.assertEqual(pur[1][4:], [500.0, 30.0, 30.0, 60.0, 560.0])
        self.assertEqual(rows(wb, 'Purchases by GST rate')[1][:2], [12, 500.0])

    def test_debit_note(self):
        dn = PurchaseReturn.objects.create(purchase=self.purchase, date=today())
        PurchaseReturnItem.objects.create(purchase_return=dn, purchase_item=self.pitem, qty=1)
        sheet = rows(self.book(), 'Debit Notes')
        self.assertTrue(sheet[1][1].startswith('DN/'))
        self.assertEqual(sheet[1][6], 100.0)

    def test_summary_numbers(self):
        s = ca_export.summary(self.start, self.end)
        self.assertEqual(s['sales']['total'], Decimal('112.00'))
        self.assertEqual(s['credit']['total'], Decimal('44.80'))
        self.assertEqual(s['purchases']['total'], Decimal('560.00'))
        self.assertEqual((s['bills'], s['credit_notes'], s['purchase_count']), (1, 1, 1))
        summary = {r[0]: r for r in rows(self.book(), 'Summary')}
        self.assertEqual(summary['Net sales'][1], 60.0)
        self.assertEqual(summary['Net sales'][5], 67.2)
        self.assertEqual(summary['GST payable = output GST - input GST'][5], round(12 - 4.8 - 60, 2))

    def test_cancelled_bills_are_not_counted(self):
        bill, _ = self.sell(qty=5, rate='10.00')
        bill.cancel()
        s = ca_export.summary(self.start, self.end)
        self.assertEqual((s['bills'], s['cancelled']), (1, 1))
        self.assertEqual(s['sales']['total'], Decimal('112.00'))

    def test_only_the_chosen_dates_are_included(self):
        wb = self.book(today() - timedelta(days=30), today() - timedelta(days=20))
        self.assertEqual(len(rows(wb, 'Sales Register')), 1)               # headers only
        self.assertEqual(len(rows(wb, 'Purchase Register')), 1)

    def test_expenses_payments_and_receipts(self):
        Expense.objects.create(name='RENT', amount=Decimal('1000'), date=today())
        cust = Customer.objects.create(name='Ramesh')
        from backend.models_extra import CustomerPayment
        CustomerPayment.objects.create(customer=cust, amount=Decimal('50'), date=today())
        wb = self.book()
        self.assertEqual(rows(wb, 'Expenses')[1][1], 'Rent')
        self.assertEqual(rows(wb, 'Expenses')[1][4], 1000.0)
        self.assertEqual(rows(wb, 'Customer Receipts')[1][1:3], ['Ramesh', 'Cash'])

    def test_bad_ranges_are_refused(self):
        with self.assertRaises(ca_export.RangeError):
            ca_export.build_workbook(today(), today() - timedelta(days=1))
        with self.assertRaises(ca_export.RangeError):
            ca_export.build_workbook(today() - timedelta(days=900), today())

    def test_page_and_download(self):
        resp = self.client.get(reverse('ca_export'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Download Excel')
        resp = self.client.get(reverse('ca_export'), {'from': self.start.isoformat(), 'to': self.end.isoformat(), 'export': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertEqual(load_workbook(io.BytesIO(resp.content)).sheetnames[0], 'Summary')
        bad = self.client.get(reverse('ca_export'), {'from': '2026-10-10', 'to': '2026-01-01', 'export': 'xlsx'})
        self.assertEqual(bad.status_code, 200)
        self.assertContains(bad, 'after the')
        junk = self.client.get(reverse('ca_export'), {'from': 'abc', 'to': 'xyz'})
        self.assertEqual(junk.status_code, 200)

    def test_only_the_owner_can_open_it(self):
        clerk = get_user_model().objects.create_user('clerk', password='pw-12345')
        self.client.force_login(clerk)
        self.assertEqual(self.client.get(reverse('ca_export')).status_code, 403)
        self.assertEqual(self.client.get(reverse('ca_export'), {'export': 'xlsx'}).status_code, 403)
