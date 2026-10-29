"""Supplier payments: a note, a date that can be chosen (not in the future, not in a closed period) and the UPI / Cheque modes."""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings

from backend.models import Payment, ShopProfile, Supplier

from .helpers import StockMixin, today

SAFE = dict(SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False, CSRF_COOKIE_SECURE=False, ALLOWED_HOSTS=['testserver'],
            MULTI_SHOP=False)


@override_settings(**SAFE)
class PaymentNoteTests(StockMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.owner = get_user_model().objects.create_superuser('boss', 'b@example.com', 'pw-12345-xx')
        self.client.force_login(self.owner)
        self.buy(qty=1, price='1000.00', invoice='INV-1')          # the supplier is owed 1000

    def owed(self):
        return Supplier.objects.get(pk=self.supplier.pk).opening_balance

    def post(self, **extra):
        data = {'supplier': self.supplier.pk, 'amount': '400', 'payment_mode': 'UPI', 'payment_date': str(today()), 'note': ''}
        data.update(extra)
        return self.client.post('/payments/new/', data, follow=True)

    def close_books(self, day):
        shop = ShopProfile.get()
        shop.books_closed_until = day
        shop.save()

    # ---- the form ----
    def test_the_form_has_the_date_and_the_note(self):
        html = self.client.get('/payments/new/').content.decode()
        self.assertIn('name="payment_date"', html)
        self.assertIn('name="note"', html)
        self.assertIn(f'value="{today():%Y-%m-%d}"', html)                 # today is filled in
        self.assertIn('UPI', html)
        self.assertIn('Cheque', html)

    def test_a_payment_with_a_note_a_date_and_upi(self):
        day = today() - timedelta(days=3)
        self.post(payment_date=str(day), note='UTR 4231987 for INV-1', payment_mode='UPI')
        pay = Payment.objects.get()
        self.assertEqual((pay.payment_date, pay.note, pay.payment_mode), (day, 'UTR 4231987 for INV-1', 'UPI'))
        self.assertEqual(self.owed(), Decimal('600.00'))                   # the balance works as before
        self.assertEqual(pay.due, Decimal('600.00'))

    def test_a_cheque_payment(self):
        self.post(payment_mode='CHEQUE', note='cheque 004512')
        self.assertEqual(Payment.objects.get().get_payment_mode_display(), 'Cheque')

    def test_the_note_is_optional_and_trimmed(self):
        self.post(note='   ')
        self.assertEqual(Payment.objects.get().note, '')
        self.post(note='  hello  ')
        self.assertEqual(Payment.objects.order_by('-id').first().note, 'hello')

    def test_a_missing_or_empty_date_means_today(self):
        self.client.post('/payments/new/', {'supplier': self.supplier.pk, 'amount': '100', 'payment_mode': 'CASH'})     # an old form without the date
        self.post(payment_date='', amount='100')
        self.assertEqual([p.payment_date for p in Payment.objects.all()], [today(), today()])

    def test_the_note_has_a_length_limit(self):
        self.post(note='x' * 300)
        self.assertEqual(Payment.objects.count(), 0)

    def test_a_future_date_is_refused(self):
        resp = self.post(payment_date=str(today() + timedelta(days=1)))
        self.assertContains(resp, 'cannot be in the future')
        self.assertEqual(Payment.objects.count(), 0)
        self.assertEqual(self.owed(), Decimal('1000.00'))

    def test_a_bad_date_is_refused_without_crashing(self):
        for bad in ('nope', '2026-13-45', '99999-01-01'):
            self.assertEqual(self.post(payment_date=bad).status_code, 200)
        self.assertEqual(Payment.objects.count(), 0)

    # ---- the pages that show it ----
    def test_the_list_shows_mode_date_and_note_and_can_search_the_note(self):
        self.post(note='UTR 99887766')
        html = self.client.get('/payments/').content.decode()
        self.assertIn('UTR 99887766', html)
        self.assertIn('UPI', html)
        self.assertIn('Note', html)
        self.assertIn('UTR 99887766', self.client.get('/payments/', {'q': '99887766'}).content.decode())
        self.assertNotIn('UTR 99887766', self.client.get('/payments/', {'q': 'zzz-nothing'}).content.decode())

    def test_the_ledger_shows_the_note_and_a_back_dated_payment_in_order(self):
        old = today() - timedelta(days=10)
        self.post(payment_date=str(old), note='cheque 77')
        html = self.client.get(f'/suppliers/{self.supplier.pk}/ledger/').content.decode()
        self.assertIn('Payment (UPI)', html)
        self.assertIn('cheque 77', html)
        from backend.views_extra import _supplier_statement
        rows, balance = _supplier_statement(self.supplier)
        self.assertEqual(balance, self.owed())                                # still equal to the stored balance
        self.assertEqual([r['date'] for r in rows], sorted(r['date'] for r in rows))

    def test_the_list_can_be_exported(self):
        self.post(note='=1+1')
        for kind in ('csv', 'xlsx'):
            self.assertEqual(self.client.get('/payments/', {'export': kind}).status_code, 200)

    def test_the_accountants_file_has_the_note_column_and_cannot_run_a_formula(self):
        from openpyxl import load_workbook
        from backend import ca_export
        self.post(note='=1+1')
        wb = load_workbook(BytesIO(ca_export.build_workbook(today() - timedelta(days=30), today())))
        ws = wb['Supplier Payments']
        self.assertEqual([c.value for c in ws[1]], ['Date', 'Supplier', 'Mode', 'Note', 'Amount'])
        self.assertEqual(ws.cell(row=2, column=4).value, "'=1+1")

    # ---- editing, deleting, closed periods ----
    def test_a_payment_can_be_edited_and_the_balance_follows(self):
        self.post(amount='400')
        pay = Payment.objects.get()
        resp = self.client.post(f'/payments/{pay.pk}/', {'supplier': self.supplier.pk, 'amount': '300', 'payment_mode': 'BANK',
                                                        'payment_date': str(today()), 'note': 'corrected'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        pay.refresh_from_db()
        self.assertEqual((pay.amount, pay.note, pay.payment_mode), (Decimal('300.00'), 'corrected', 'BANK'))
        self.assertEqual(self.owed(), Decimal('700.00'))

    def test_changing_only_the_note_does_not_touch_the_balance(self):
        self.post(amount='400')
        pay = Payment.objects.get()
        pay.note = 'now with a note'
        pay.save()
        self.assertEqual(self.owed(), Decimal('600.00'))

    def test_a_payment_in_a_closed_period_cannot_be_added(self):
        self.close_books(today() - timedelta(days=5))
        resp = self.post(payment_date=str(today() - timedelta(days=9)))
        self.assertContains(resp, 'closed period')
        self.assertEqual(Payment.objects.count(), 0)
        self.post(payment_date=str(today()))                                  # an open day still works
        self.assertEqual(Payment.objects.count(), 1)

    def test_a_payment_of_a_closed_period_keeps_its_figures_but_its_note_can_be_corrected(self):
        old = today() - timedelta(days=9)
        pay = Payment.objects.create(supplier=self.supplier, amount=Decimal('100'), payment_mode='CASH', payment_date=old)
        self.close_books(today() - timedelta(days=5))
        pay.note = 'UTR added later'
        pay.save()                                                            # allowed: the figures did not change
        pay.refresh_from_db()
        self.assertEqual(pay.note, 'UTR added later')
        pay.amount = Decimal('150')
        with self.assertRaises(ValidationError):
            pay.save()
        pay.refresh_from_db()
        self.assertEqual(pay.amount, Decimal('100.00'))
        with self.assertRaises(ValidationError):
            pay.delete()
        self.assertTrue(Payment.objects.filter(pk=pay.pk).exists())

    def test_deleting_such_a_payment_from_the_page_shows_a_message_not_an_error(self):
        pay = Payment.objects.create(supplier=self.supplier, amount=Decimal('100'), payment_mode='CASH', payment_date=today() - timedelta(days=9))
        self.close_books(today() - timedelta(days=5))
        resp = self.client.post(f'/payments/{pay.pk}/delete/', follow=True)
        self.assertContains(resp, 'closed period')
        self.assertTrue(Payment.objects.filter(pk=pay.pk).exists())

    def test_an_open_payment_can_still_be_deleted_and_the_balance_returns(self):
        self.post(amount='400')
        pay = Payment.objects.get()
        self.client.post(f'/payments/{pay.pk}/delete/', follow=True)
        self.assertEqual(Payment.objects.count(), 0)
        self.assertEqual(self.owed(), Decimal('1000.00'))

    def test_the_admin_page_opens(self):
        pay = Payment.objects.create(supplier=self.supplier, amount=Decimal('100'), payment_mode='CASH')
        self.assertEqual(self.client.get('/admin/backend/payment/').status_code, 200)
        self.assertEqual(self.client.get(f'/admin/backend/payment/{pay.pk}/change/').status_code, 200)
