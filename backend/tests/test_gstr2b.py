"""GSTR-2B check: reading the portal's JSON, matching with the purchases, and the page."""
import io
import json
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import load_workbook

from backend import gstr2b
from backend.models import Supplier

from .helpers import StockMixin, today

G1 = '27AAPFU0939F1ZV'
G2 = '07AAGFF2194N1Z1'
FIRST = date(2026, 9, 1)
LAST = date(2026, 9, 30)


def inv(num, taxable=500, cgst=30, sgst=30, igst=0, day='05-09-2026', **kw):
    d = dict(inum=num, dt=day, val=taxable + cgst + sgst + igst, txval=taxable, cgst=cgst, sgst=sgst, igst=igst, cess=0,
             itcavl='Y', rev='N')
    d.update(kw)
    return d


def portal(*suppliers, **extra):
    """A GSTR-2B JSON like the portal's. suppliers: (ctin, name, [invoices])."""
    doc = {'b2b': [dict(ctin=c, trdnm=n, inv=i) for c, n, i in suppliers]}
    doc.update(extra)
    return json.dumps({'data': {'gstin': '24AAACC1206D1ZM', 'rtnprd': '092026', 'docdata': doc}}).encode()


def book(number, gstin=G1, day=date(2026, 9, 5), taxable=500, tax=60, supplier='ABC Pharma'):
    d = Decimal
    return dict(id=1, number=number, date=day, gstin=gstin, supplier=supplier, taxable=d(taxable), tax=d(tax), total=d(taxable + tax))


def run(json_bytes, books):
    return gstr2b.reconcile(gstr2b.parse_json(json_bytes)['docs'], books, FIRST, LAST)


def statuses(result):
    return [r['status'] for r in result['rows']]


class NumberTests(SimpleTestCase):
    def test_invoice_numbers_are_compared_loosely(self):
        n = gstr2b.norm_number
        self.assertEqual(n('INV/2026/001'), n('inv-2026-1'))
        self.assertEqual(n(' 000123 '), n('123'))
        self.assertEqual(n('A-100'), 'A100')
        self.assertNotEqual(n('A-100'), n('A-1'))
        self.assertNotEqual(n('INV-1'), n('INV-2'))

    def test_dates(self):
        self.assertEqual(gstr2b.parse_date('05-09-2026'), date(2026, 9, 5))
        self.assertEqual(gstr2b.parse_date('2026-09-05'), date(2026, 9, 5))
        self.assertIsNone(gstr2b.parse_date('nonsense'))


class ParseJsonTests(SimpleTestCase):
    def test_invoices_and_notes_are_read(self):
        raw = portal((G1, 'ABC Pharma', [inv('INV-1')]),
                     cdnr=[dict(ctin=G1, nt=[dict(ntnum='CN-1', typ='C', dt='10-09-2026', val=112, txval=100, cgst=6, sgst=6, igst=0)])])
        p = gstr2b.parse_json(raw)
        self.assertEqual(p['period'], FIRST)
        self.assertEqual(p['gstin'], '24AAACC1206D1ZM')
        kinds = sorted(d['kind'] for d in p['docs'])
        self.assertEqual(kinds, ['CN', 'INV'])
        d = next(x for x in p['docs'] if x['kind'] == 'INV')
        self.assertEqual((d['taxable'], d['tax'], d['gstin']), (Decimal('500'), Decimal('60'), G1))

    def test_amounts_given_per_item_are_added_up(self):
        node = dict(inum='X-1', dt='01-09-2026', val=1120,
                    items=[dict(txval=400, igst=48), dict(txval=600, igst=72)])
        p = gstr2b.parse_json(portal((G1, 'ABC', [node])))
        d = p['docs'][0]
        self.assertEqual((d['taxable'], d['igst'], d['tax']), (Decimal('1000'), Decimal('120'), Decimal('120')))

    def test_itc_not_available_and_reverse_charge_are_flagged(self):
        p = gstr2b.parse_json(portal((G1, 'ABC', [inv('X-1', itcavl='N', rsn='POS'), inv('X-2', rev='Y')])))
        a, b = p['docs']
        self.assertFalse(a['itc'])
        self.assertIn('not available', a['note'])
        self.assertIn('Reverse charge', b['note'])

    def test_a_wrong_file_gives_a_clear_message(self):
        for raw in (b'not json at all', b'[]', b'{"data": {"x": 1}}', json.dumps({'hello': 1}).encode()):
            with self.subTest(raw=raw[:20]):
                with self.assertRaises(gstr2b.FileProblem):
                    gstr2b.parse_json(raw)

    def test_sections_that_are_not_compared_are_reported(self):
        p = gstr2b.parse_json(portal((G1, 'ABC', [inv('X-1')]), b2ba=[{'x': 1}], impg=[{'y': 1}]))
        self.assertEqual(p['skipped'], ['b2ba', 'impg'])


class ParseRowsTests(SimpleTestCase):
    def test_a_plain_list_with_usual_headers(self):
        headers = ['gstin_of_supplier', 'invoice_number', 'invoice_date', 'taxable_value', 'central_tax', 'state_ut_tax', 'integrated_tax']
        rows = [(2, dict(zip(headers, [G1, 'INV-1', '05/09/2026', '500', '30', '30', '0'])))]
        p = gstr2b.parse_rows(headers, rows)
        d = p['docs'][0]
        self.assertEqual((d['number'], d['taxable'], d['tax'], d['date']), ('INV-1', Decimal('500'), Decimal('60'), date(2026, 9, 5)))

    def test_missing_columns_are_named(self):
        with self.assertRaises(gstr2b.FileProblem) as ctx:
            gstr2b.parse_rows(['foo', 'bar'], [(2, {'foo': 1, 'bar': 2})])
        self.assertIn('Invoice number', str(ctx.exception))
        self.assertIn('GSTIN', str(ctx.exception))

    def test_rows_without_an_invoice_number_are_skipped_and_an_empty_file_is_refused(self):
        headers = ['gstin', 'invoice_no', 'taxable']
        with self.assertRaises(gstr2b.FileProblem):
            gstr2b.parse_rows(headers, [(2, {'gstin': G1, 'invoice_no': '', 'taxable': '5'})])


class MatchingTests(SimpleTestCase):
    def test_same_bill_is_matched(self):
        r = run(portal((G1, 'ABC', [inv('INV-1')])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.MATCHED])
        self.assertEqual(r['totals']['itc_matched'], Decimal('60'))

    def test_a_rounding_difference_is_still_a_match_but_a_real_difference_is_not(self):
        r = run(portal((G1, 'ABC', [inv('INV-1', taxable=500, cgst=30.4, sgst=30.4)])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.MATCHED])
        r = run(portal((G1, 'ABC', [inv('INV-1', taxable=500, cgst=15, sgst=15)])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.MISMATCH])
        self.assertEqual(r['rows'][0]['diff_tax'], Decimal('-30'))
        self.assertEqual(r['totals']['itc_at_risk'], Decimal('30'))      # you booked 30 more tax than the supplier filed

    def test_different_spelling_of_the_number_is_the_same_bill(self):
        r = run(portal((G1, 'ABC', [inv('inv/001')])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.MATCHED])

    def test_bill_only_in_your_books_means_itc_at_risk(self):
        r = run(portal((G1, 'ABC', [])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.ONLY_BOOKS])
        self.assertEqual(r['totals']['itc_at_risk'], Decimal('60'))

    def test_bill_only_in_2b_is_not_in_your_books(self):
        r = run(portal((G1, 'ABC', [inv('INV-9')])), [])
        self.assertEqual(statuses(r), [gstr2b.ONLY_2B])
        self.assertEqual(r['totals']['itc_not_in_books'], Decimal('60'))

    def test_a_typing_mistake_in_the_invoice_number_is_suggested(self):
        r = run(portal((G1, 'ABC', [inv('IN-1')])), [book('INV-1')])
        self.assertEqual(statuses(r), [gstr2b.NUMBER_DIFF])
        self.assertIn('INV-1', r['rows'][0]['note'])

    def test_the_right_number_with_a_wrong_supplier_gstin_is_suggested(self):
        r = run(portal((G2, 'Other', [inv('INV-1')])), [book('INV-1', gstin=G1)])
        self.assertEqual(statuses(r), [gstr2b.GSTIN_DIFF])

    def test_purchases_from_suppliers_without_gstin_are_counted_apart(self):
        r = run(portal((G1, 'ABC', [])), [book('INV-1', gstin='')])
        self.assertEqual(statuses(r), [])
        self.assertEqual(r['totals']['no_gstin'], 1)

    def test_an_earlier_invoice_in_this_months_2b_is_matched_but_old_books_are_not_reported(self):
        old = book('INV-7', day=date(2026, 8, 20))
        r = run(portal((G1, 'ABC', [inv('INV-7', day='20-08-2026')])), [old])
        self.assertEqual(statuses(r), [gstr2b.MATCHED])
        r = run(portal((G1, 'ABC', [])), [old])               # last month's bill is not "missing" from this month's 2B
        self.assertEqual(statuses(r), [])

    def test_a_date_in_another_month_is_noted(self):
        r = run(portal((G1, 'ABC', [inv('INV-1', day='28-08-2026')])), [book('INV-1', day=date(2026, 9, 2))])
        self.assertEqual(statuses(r), [gstr2b.MATCHED])
        self.assertIn('Dates differ', r['rows'][0]['note'])

    def test_itc_blocked_invoices_and_notes_are_totalled(self):
        raw = portal((G1, 'ABC', [inv('INV-1', itcavl='N')]),
                     cdnr=[dict(ctin=G1, nt=[dict(ntnum='C1', typ='C', dt='10-09-2026', val=56, txval=50, cgst=3, sgst=3)])])
        r = run(raw, [book('INV-1')])
        self.assertEqual(r['totals']['itc_blocked'], Decimal('60'))
        self.assertEqual((r['totals']['credit_note_count'], r['totals']['credit_notes']), (1, Decimal('6')))

    def test_every_bill_is_used_once(self):
        raw = portal((G1, 'ABC', [inv('INV-1'), inv('INV-1')]))
        r = run(raw, [book('INV-1')])
        self.assertEqual(sorted(statuses(r)), sorted([gstr2b.MATCHED, gstr2b.ONLY_2B]))

    def test_the_problem_rows_come_first(self):
        raw = portal((G1, 'ABC', [inv('INV-1'), inv('INV-2', cgst=1, sgst=1)]))
        r = run(raw, [book('INV-1'), book('INV-2'), book('INV-3')])
        self.assertEqual(statuses(r), [gstr2b.MISMATCH, gstr2b.ONLY_BOOKS, gstr2b.MATCHED])


@override_settings(SECURE_SSL_REDIRECT=False, SESSION_COOKIE_SECURE=False, CSRF_COOKIE_SECURE=False)
class GstrPageTests(StockMixin, TestCase):
    def setUp(self):
        super().setUp()
        Supplier.objects.filter(pk=self.supplier.pk).update(gstin=G1)
        self.owner = get_user_model().objects.create_superuser('owner', 'o@example.com', 'pw-12345')
        self.client.force_login(self.owner)
        self.buy(qty=5, price='560.00', invoice='INV-1')                   # taxable 500, GST 60, dated today
        self.month = f"{today():%Y-%m}"
        self.first = today().replace(day=1)

    def post(self, raw, name='2b.json', **data):
        data.setdefault('month', self.month)
        return self.client.post('/gstr2b/', dict(data, file=SimpleUploadedFile(name, raw)))

    def payload(self, *invoices):
        day = f"{today():%d-%m-%Y}"
        invs = [inv(n, day=day, **kw) for n, kw in invoices]
        return portal((G1, 'ABC Pharma', invs))

    def test_the_books_are_read_with_taxable_value_and_gst(self):
        rows = gstr2b.books_between(self.first, today())
        self.assertEqual(len(rows), 1)
        b = rows[0]
        self.assertEqual((b['number'], b['gstin'], b['taxable'], b['tax'], b['total']),
                         ('INV-1', G1, Decimal('500.00'), Decimal('60.00'), Decimal('560.00')))

    def test_the_page_opens_for_the_owner_only(self):
        self.assertEqual(self.client.get('/gstr2b/').status_code, 200)
        clerk = get_user_model().objects.create_user('clerk', password='pw-12345')
        self.client.force_login(clerk)
        self.assertEqual(self.client.get('/gstr2b/').status_code, 403)

    def test_a_matching_file_shows_matched(self):
        r = self.post(self.payload(('INV-1', {})))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['result']['totals']['counts']['MATCHED'], 1)
        self.assertContains(r, 'Matched invoices')

    def test_a_missing_bill_is_reported(self):
        r = self.post(self.payload(('INV-OTHER', dict(taxable=100, cgst=6, sgst=6))))
        counts = r.context['result']['totals']['counts']
        self.assertEqual((counts['ONLY_BOOKS'], counts['ONLY_2B']), (1, 1))

    def test_a_bad_file_gives_a_message_and_no_error_page(self):
        r = self.post(b'{broken')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'not a valid JSON')
        r = self.post(b'x', name='list.pdf')
        self.assertContains(r, 'Please upload')
        r = self.post(json.dumps({'data': {'docdata': {'b2b': [5]}}}).encode())
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'could not be read')

    def test_no_file_is_a_message(self):
        r = self.client.post('/gstr2b/', {'month': self.month})
        self.assertContains(r, 'choose the GSTR-2B file')

    def test_the_month_in_the_file_is_used_when_none_is_chosen(self):
        r = self.post(portal((G1, 'ABC', [])), month='')
        self.assertEqual(r.context['first'], date(2026, 9, 1))

    def test_excel_download(self):
        r = self.post(self.payload(('INV-1', {})), export='1')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Summary', 'Reconciliation'])
        self.assertEqual(wb['Reconciliation'].max_row, 2)

    def test_a_csv_list_works_too(self):
        csv = f"GSTIN,Invoice number,Invoice date,Taxable value,Central tax,State tax\n{G1},INV-1,{today():%d/%m/%Y},500,30,30\n"
        r = self.post(csv.encode(), name='list.csv')
        self.assertEqual(r.context['result']['totals']['counts']['MATCHED'], 1)

    def test_a_gstin_that_is_not_the_shops_gets_a_warning(self):
        from backend.models import ShopProfile
        shop = ShopProfile.get()
        shop.gstin = '27AAPFU0939F1ZV'
        shop.save()
        r = self.post(self.payload(('INV-1', {})))
        self.assertTrue(any('Shop Settings' in w for w in r.context['warn']))
