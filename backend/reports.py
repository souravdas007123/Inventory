"""
Report calculations in one place, so the web page, admin and Excel export all show the same numbers.

- profit_and_loss      : P&L (returns and write-offs included)
- gst_report           : GSTR-1 (B2B/B2C/Credit notes), HSN summary, GSTR-3B summary
- drug_register        : Schedule H/H1/X sales register
- dashboard_charts     : daily sales, top selling, dead stock
- gst_workbook / register_workbook : Excel (.xlsx) export
"""
import io
from collections import OrderedDict
from datetime import timedelta
from decimal import Decimal

from django.db.models import Count, F, Min, Sum
from django.db.models.functions import ExtractWeekDay, TruncMonth
from django.utils import timezone

from . import supplier_prices
from .notify import credit_reminder_link
from .models import (
    Batch, Bill, Customer, CustomerPayment, Expense, InvoiceItem, Product, PurchaseItem,
    PurchaseReturnItem, Sale, SalesReturn, SalesReturnItem, ShopProfile, StockWriteOff, Supplier, Transaction,
)

ZERO = Decimal('0')
TWO = Decimal('0.01')


def dec(x):
    return Decimal(str(x or 0))


def q2(x):
    return Decimal(x).quantize(TWO)


# ======================================================
# PROFIT & LOSS
# ======================================================

def profit_and_loss(month, year):
    """
    Income    = taxable value of Net Sales (after returns; GST is not income)
    COGS      = purchase cost of goods sold (from batch_history) - cost of returned goods
    Write-off = purchase cost of expired/damaged goods (loss)
    Purchases = for information only (taxable goods purchased this month, after returns)
    """
    # Returns are stored as negative Sale rows, so this sum is net
    sales = dec(Sale.objects.filter(date__month=month, date__year=year)
                .aggregate(s=Sum('taxable_value'))['s'])

    histories = list(InvoiceItem.objects.filter(sale__date__month=month, sale__date__year=year)
                     .values_list('batch_history', flat=True))
    batch_ids = {int(k) for h in histories for k in (h or {})}
    batch_rate = dict(Batch.objects.filter(pk__in=batch_ids, purchase_item__isnull=False)
                      .values_list('pk', 'purchase_item__rate'))
    cogs = ZERO
    for history in histories:
        for batch_id, qty in (history or {}).items():
            cogs += dec(batch_rate.get(int(batch_id))) * qty
    cogs -= dec(SalesReturnItem.objects.filter(
        sales_return__date__month=month, sales_return__date__year=year
    ).aggregate(s=Sum('cost'))['s'])

    purchases = dec(PurchaseItem.objects.filter(
        purchase__invoice_date__month=month, purchase__invoice_date__year=year
    ).aggregate(s=Sum('tax'))['s'])
    purchases -= dec(PurchaseReturnItem.objects.filter(
        purchase_return__date__month=month, purchase_return__date__year=year
    ).aggregate(s=Sum('taxable_value'))['s'])

    expenses = dec(Expense.objects.filter(date__month=month, date__year=year)
                   .aggregate(s=Sum('amount'))['s'])
    writeoffs = dec(StockWriteOff.objects.filter(date__month=month, date__year=year)
                    .aggregate(s=Sum('cost'))['s'])

    gross = sales - cogs
    return dict(sales=sales, cogs=cogs, gross=gross, purchases=purchases,
                expenses=expenses, writeoffs=writeoffs, net=gross - expenses - writeoffs)


# ======================================================
# GST REPORTS
# ======================================================

def _tax_row():
    return dict(taxable=ZERO, cgst=ZERO, sgst=ZERO, igst=ZERO, total=ZERO)


def _add(row, taxable, cgst, sgst, igst, total, sign=1):
    row['taxable'] += sign * dec(taxable)
    row['cgst'] += sign * dec(cgst)
    row['sgst'] += sign * dec(sgst)
    row['igst'] += sign * dec(igst)
    row['total'] += sign * dec(total)


def gst_report(month, year):
    """
    Monthly summary similar to GSTR-1 / GSTR-3B. Only FINAL bills and that month's credit notes are counted.
    Note: this is an accounting summary; verify it with your CA before filing a return.
    """
    items = (InvoiceItem.objects
             .filter(bill__status='FINAL', bill__date__month=month, bill__date__year=year)
             .select_related('bill', 'bill__customer', 'product'))

    b2b = OrderedDict()      # bill id -> row
    b2c = OrderedDict()      # gst rate -> row
    hsn = OrderedDict()      # (hsn, rate) -> row
    outward = _tax_row()

    for it in sorted(items, key=lambda i: (i.bill.date, i.bill_id, i.id)):
        bill = it.bill
        gstin = bill.customer.gstin if bill.customer_id and bill.customer.gstin else ''
        if gstin:
            row = b2b.setdefault(bill.pk, dict(
                invoice=bill.display_number, date=bill.date, customer=bill.customer_name,
                gstin=gstin, **_tax_row()))
        else:
            row = b2c.setdefault(it.gst_rate, dict(rate=it.gst_rate, **_tax_row()))
        _add(row, it.taxable_value, it.cgst, it.sgst, it.igst, it.total)

        key = ((it.product.hsn_code if it.product_id and it.product.hsn_code else '-'), it.gst_rate)
        hrow = hsn.setdefault(key, dict(hsn=key[0], rate=key[1], qty=0, **_tax_row()))
        hrow['qty'] += it.base_qty
        _add(hrow, it.taxable_value, it.cgst, it.sgst, it.igst, it.total)
        _add(outward, it.taxable_value, it.cgst, it.sgst, it.igst, it.total)

    # ---- Credit notes (sales return) ----
    credit_notes = OrderedDict()
    returns = (SalesReturnItem.objects
               .filter(sales_return__date__month=month, sales_return__date__year=year)
               .select_related('sales_return', 'sales_return__bill', 'sales_return__bill__customer',
                               'invoice_item', 'invoice_item__product')
               .order_by('sales_return__date', 'sales_return_id', 'id'))
    for r in returns:
        sr = r.sales_return
        bill = sr.bill
        gstin = bill.customer.gstin if bill.customer_id and bill.customer.gstin else ''
        row = credit_notes.setdefault(sr.pk, dict(
            note=sr.note_number, date=sr.date, invoice=bill.display_number,
            customer=bill.customer_name, gstin=gstin, **_tax_row()))
        _add(row, r.taxable_value, r.cgst, r.sgst, r.igst, r.total)
        _add(outward, r.taxable_value, r.cgst, r.sgst, r.igst, r.total, sign=-1)
        ii = r.invoice_item
        key = ((ii.product.hsn_code if ii.product_id and ii.product.hsn_code else '-'), ii.gst_rate)
        hrow = hsn.setdefault(key, dict(hsn=key[0], rate=key[1], qty=0, **_tax_row()))
        hrow['qty'] -= r.base_qty
        _add(hrow, r.taxable_value, r.cgst, r.sgst, r.igst, r.total, sign=-1)
        # B2C return: reduce the net rate-wise
        if not gstin:
            crow = b2c.setdefault(ii.gst_rate, dict(rate=ii.gst_rate, **_tax_row()))
            _add(crow, r.taxable_value, r.cgst, r.sgst, r.igst, r.total, sign=-1)

    for row in credit_notes.values():
        row['gst'] = row['cgst'] + row['sgst'] + row['igst']

    # ---- Documents issued ----
    bills = Bill.objects.filter(date__month=month, date__year=year).order_by('invoice_number')
    numbers = [b.invoice_number for b in bills if b.invoice_number]
    docs = dict(
        invoices=len(numbers), first=numbers[0] if numbers else '-', last=numbers[-1] if numbers else '-',
        cancelled=sum(1 for b in bills if b.status == 'CANCELLED'), credit_notes=len(credit_notes),
    )

    # ---- ITC (purchases) ----
    pur = PurchaseItem.objects.filter(purchase__invoice_date__month=month, purchase__invoice_date__year=year)
    itc = dict(taxable=dec(pur.aggregate(s=Sum('tax'))['s']), cgst=dec(pur.aggregate(s=Sum('cgst'))['s']),
               sgst=dec(pur.aggregate(s=Sum('sgst'))['s']), igst=ZERO)
    for r in PurchaseReturnItem.objects.filter(purchase_return__date__month=month,
                                               purchase_return__date__year=year):
        half = q2(r.gst / 2)
        itc['taxable'] -= r.taxable_value
        itc['cgst'] -= half
        itc['sgst'] -= (r.gst - half)

    net = dict(cgst=outward['cgst'] - itc['cgst'], sgst=outward['sgst'] - itc['sgst'],
               igst=outward['igst'] - itc['igst'])
    net['total'] = net['cgst'] + net['sgst'] + net['igst']

    return dict(
        month=month, year=year,
        b2b=list(b2b.values()), b2c=sorted(b2c.values(), key=lambda r: r['rate']),
        credit_notes=list(credit_notes.values()),
        hsn=sorted(hsn.values(), key=lambda r: (str(r['hsn']), r['rate'])),
        docs=docs, outward=outward, itc=itc, net=net,
    )


def _sheet(wb, title, headers, rows, first=False):
    from openpyxl.styles import Font
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for idx, header in enumerate(headers, 1):
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows if len(r) >= idx] or [10])
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(width + 2, 10), 40)
    return ws


def _n(x):
    return float(x)


def gst_workbook(rep):
    """Converts the GST report to Excel (.xlsx) bytes."""
    from openpyxl import Workbook
    wb = Workbook()
    tax_h = ['Taxable', 'CGST', 'SGST', 'IGST', 'Total']

    def tax_cells(r):
        return [_n(r['taxable']), _n(r['cgst']), _n(r['sgst']), _n(r['igst']), _n(r['total'])]

    _sheet(wb, 'GSTR-3B Summary', ['Particular', 'Taxable', 'CGST', 'SGST', 'IGST'], [
        ['Outward supplies (net of credit notes)', _n(rep['outward']['taxable']), _n(rep['outward']['cgst']),
         _n(rep['outward']['sgst']), _n(rep['outward']['igst'])],
        ['ITC from purchases (net of debit notes)', _n(rep['itc']['taxable']), _n(rep['itc']['cgst']),
         _n(rep['itc']['sgst']), _n(rep['itc']['igst'])],
        ['Net GST payable', '', _n(rep['net']['cgst']), _n(rep['net']['sgst']), _n(rep['net']['igst'])],
    ], first=True)
    _sheet(wb, 'B2B Invoices', ['Invoice', 'Date', 'Customer', 'GSTIN'] + tax_h,
           [[r['invoice'], r['date'].strftime('%d-%m-%Y'), r['customer'], r['gstin']] + tax_cells(r)
            for r in rep['b2b']])
    _sheet(wb, 'B2C (rate wise)', ['GST %'] + tax_h, [[r['rate']] + tax_cells(r) for r in rep['b2c']])
    _sheet(wb, 'Credit Notes', ['Note', 'Date', 'Against Invoice', 'Customer', 'GSTIN'] + tax_h,
           [[r['note'], r['date'].strftime('%d-%m-%Y'), r['invoice'], r['customer'], r['gstin']] + tax_cells(r)
            for r in rep['credit_notes']])
    _sheet(wb, 'HSN Summary', ['HSN', 'GST %', 'Qty (base unit)'] + tax_h,
           [[r['hsn'], r['rate'], r['qty']] + tax_cells(r) for r in rep['hsn']])
    d = rep['docs']
    _sheet(wb, 'Documents', ['Invoices issued', 'First', 'Last', 'Cancelled', 'Credit notes'],
           [[d['invoices'], d['first'], d['last'], d['cancelled'], d['credit_notes']]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================================================
# DRUG REGISTER (Schedule H / H1 / X)
# ======================================================

def drug_register(month, year, schedules=('H1', 'X')):
    items = list(InvoiceItem.objects
                 .filter(bill__status='FINAL', bill__date__month=month, bill__date__year=year,
                         product__schedule__in=schedules)
                 .select_related('bill', 'product', 'unit')
                 .order_by('bill__date', 'bill_id', 'id'))
    batch_map = Batch.objects.in_bulk({int(k) for it in items for k in (it.batch_history or {})})
    rows = []
    for it in items:
        batches = []
        for key in (it.batch_history or {}):
            b = batch_map.get(int(key))
            if b:
                batches.append(f"{b.batch_number}" + (f" / {b.expire_date:%m-%y}" if b.expire_date else ''))
        bill = it.bill
        rows.append(dict(
            date=bill.date, invoice=bill.display_number, patient=bill.patient_name,
            address=bill.patient_address, doctor=bill.doctor_name, reg_no=bill.doctor_reg_no,
            product=it.product.name or str(it.product), schedule=it.product.schedule,
            batch=', '.join(batches), qty=f"{it.qty} {it.unit or ''}".strip(),
        ))
    return rows


def register_workbook(rows):
    from openpyxl import Workbook
    wb = Workbook()
    headers = ['Date', 'Invoice', 'Patient', 'Patient Address', 'Doctor', 'Doctor Reg. No.',
               'Medicine', 'Schedule', 'Batch / Exp', 'Qty']
    _sheet(wb, 'Drug Register', headers,
           [[r['date'].strftime('%d-%m-%Y'), r['invoice'], r['patient'], r['address'], r['doctor'],
             r['reg_no'], r['product'], r['schedule'], r['batch'], r['qty']] for r in rows], first=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================================================
# DASHBOARD CHARTS
# ======================================================

def dashboard_charts(today=None, days=14):
    """Daily sales (last N days), top selling (30 days), dead stock (no sales in 90 days)."""
    today = today or timezone.localdate()
    start = today - timedelta(days=days - 1)
    daily = {d: dec(s) for d, s in Sale.objects.filter(date__gte=start, date__lte=today)
             .values_list('date').annotate(s=Sum('total')).order_by()}
    series = []
    for i in range(days):
        day = start + timedelta(days=i)
        series.append((day, daily.get(day, ZERO)))
    peak = max((v for _, v in series), default=ZERO)
    bars = [dict(label=day.strftime('%d %b'), value=val,
                 pct=int(round(float(val / peak) * 100)) if peak > 0 and val > 0 else 0)
            for day, val in series]

    recent = InvoiceItem.objects.filter(bill__status='FINAL', bill__date__gte=today - timedelta(days=29))
    top = list(recent.filter(product__isnull=False).values('product_id', 'product__name')
               .annotate(amount=Sum('taxable_value'), qty=Sum('base_qty')).order_by('-amount')[:5])
    top_max = max((dec(t['amount']) for t in top), default=ZERO)
    top_rows = [dict(name=t['product__name'] or '-', amount=dec(t['amount']), qty=t['qty'],
                     pct=int(round(float(dec(t['amount']) / top_max) * 100)) if top_max > 0 else 0)
                for t in top]

    sold_ids = (InvoiceItem.objects.filter(bill__status='FINAL', bill__date__gte=today - timedelta(days=90))
                .filter(product__isnull=False).values_list('product_id', flat=True).distinct())
    dead = [(p.name or str(p), p.stock_display()) for p in
            Product.objects.exclude(is_active=False).filter(stock_qty__gt=0).exclude(pk__in=sold_ids)
            .prefetch_related('unit_conversions__unit').select_related('unit').order_by('-stock_qty')[:8]]
    return dict(bars=bars, total_14=sum((v for _, v in series), ZERO), top_rows=top_rows, dead_stock=dead)


# ======================================================
# DASHBOARD OVERVIEW (cards, lists, extra charts)
# ======================================================

def _cogs(start, end):
    """Purchase cost of goods sold between [start, end] (after returns)."""
    histories = list(InvoiceItem.objects.filter(sale__date__gte=start, sale__date__lte=end)
                     .values_list('batch_history', flat=True))
    batch_ids = {int(k) for h in histories for k in (h or {})}
    batch_rate = dict(Batch.objects.filter(pk__in=batch_ids, purchase_item__isnull=False)
                      .values_list('pk', 'purchase_item__rate'))
    cogs = ZERO
    for history in histories:
        for batch_id, qty in (history or {}).items():
            cogs += dec(batch_rate.get(int(batch_id))) * qty
    cogs -= dec(SalesReturnItem.objects.filter(
        sales_return__date__gte=start, sales_return__date__lte=end).aggregate(s=Sum('cost'))['s'])
    return cogs


def _pct(value, peak):
    return int(round(float(value / peak) * 100)) if peak and peak > 0 and value > 0 else 0


def dashboard_overview(today=None, include_money=False):
    """
    Extra dashboard numbers. With include_money=False (non-staff), profit / stock value are not computed.
    """
    today = today or timezone.localdate()
    month_start = today.replace(day=1)
    out = {}

    # ---- Bill counts ----
    final = Bill.objects.filter(status='FINAL')
    out['bills_today'] = final.filter(date=today).count()
    out['bills_month'] = final.filter(date__gte=month_start, date__lte=today).count()
    month_sale = dec(Sale.objects.filter(date__gte=month_start, date__lte=today).aggregate(s=Sum('total'))['s'])
    out['avg_bill'] = q2(month_sale / out['bills_month']) if out['bills_month'] else ZERO
    out['gst_today'] = dec(Sale.objects.filter(date=today).aggregate(s=Sum('gst'))['s'])
    out['gst_month'] = dec(Sale.objects.filter(date__gte=month_start, date__lte=today).aggregate(s=Sum('gst'))['s'])

    # ---- Credit (to collect from customers) ----
    credit_ids = Customer.objects.filter(bills__payment_mode='CREDIT', bills__status='FINAL').distinct()[:300]
    first_credit = dict(Bill.objects.filter(status='FINAL', payment_mode='CREDIT', customer__isnull=False)
                        .values_list('customer_id').annotate(d=Min('date')).order_by())
    shop_name = ShopProfile.get().name
    owing = []
    for c in credit_ids:
        bal = c.balance()
        if bal > 0:
            since = first_credit.get(c.pk)
            owing.append(dict(pk=c.pk, name=c.name, phone=c.phone or '', balance=bal,
                              days=(today - since).days if since else 0,
                              over=bool(c.credit_limit and bal > c.credit_limit),
                              wa=credit_reminder_link(shop_name, c.name, c.phone, bal)))
    owing.sort(key=lambda r: r['balance'], reverse=True)
    out['credit_total'] = sum((r['balance'] for r in owing), ZERO)
    out['credit_count'] = len(owing)
    out['credit_rows'] = owing[:6]

    # ---- Payable to suppliers ----
    owed = Supplier.objects.filter(opening_balance__gt=0)
    out['payable_total'] = dec(owed.aggregate(s=Sum('opening_balance'))['s'])
    out['payable_count'] = owed.count()
    out['payable_rows'] = [dict(name=sp.company, phone=sp.phone or '', balance=dec(sp.opening_balance))
                           for sp in owed.order_by('-opening_balance')[:6]]

    # ---- Payment mode split (this month's bills) ----
    labels = dict(Bill.PAYMENT_CHOICES)
    colors = ['var(--pri)', 'var(--amber)', '#3b82f6', 'var(--red)']
    split = list(InvoiceItem.objects.filter(bill__status='FINAL', bill__date__gte=month_start,
                                            bill__date__lte=today)
                 .values('bill__payment_mode').annotate(s=Sum('total')).order_by('-s'))
    split_total = sum((dec(r['s']) for r in split), ZERO)
    rows, acc = [], 0.0
    for i, r in enumerate(split):
        amount = dec(r['s'])
        share = float(amount / split_total * 100) if split_total > 0 else 0.0
        rows.append(dict(label=labels.get(r['bill__payment_mode'], r['bill__payment_mode']), amount=amount,
                         pct=int(round(share)), color=colors[i % len(colors)], start=acc, end=acc + share))
        acc += share
    out['pay_rows'] = rows
    out['pay_total'] = split_total
    out['pay_gradient'] = ', '.join(f"{r['color']} {r['start']:.2f}% {r['end']:.2f}%" for r in rows)

    # ---- Last 6 months: sales vs purchases ----
    months = []
    y, m = today.year, today.month
    for _ in range(6):
        months.append((y, m))
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    months.reverse()
    from datetime import date as _date
    six_start = _date(months[0][0], months[0][1], 1)
    sales_m = {(r['m'].year, r['m'].month): dec(r['s']) for r in
               Sale.objects.filter(date__gte=six_start).annotate(m=TruncMonth('date')).values('m')
               .annotate(s=Sum('total')).order_by() if r['m']}
    buy_m = {(r['m'].year, r['m'].month): dec(r['s']) for r in
             Transaction.objects.filter(purchase_date__gte=six_start).annotate(m=TruncMonth('purchase_date'))
             .values('m').annotate(s=Sum('purchase_amount')).order_by() if r['m']}
    peak = max([*sales_m.values(), *buy_m.values(), ZERO])
    out['month_bars'] = [dict(label=_date(yy, mm, 1).strftime('%b'), sale=sales_m.get((yy, mm), ZERO),
                              buy=buy_m.get((yy, mm), ZERO),
                              sale_pct=_pct(sales_m.get((yy, mm), ZERO), peak),
                              buy_pct=_pct(buy_m.get((yy, mm), ZERO), peak)) for yy, mm in months]

    # ---- Sales by weekday (last 8 weeks) ----
    names = {2: 'Mon', 3: 'Tue', 4: 'Wed', 5: 'Thu', 6: 'Fri', 7: 'Sat', 1: 'Sun'}   # ExtractWeekDay: 1=Sunday
    wk = {r['w']: dec(r['s']) for r in
          Sale.objects.filter(date__gte=today - timedelta(days=55), date__lte=today)
          .annotate(w=ExtractWeekDay('date')).values('w').annotate(s=Sum('total')).order_by()}
    wpeak = max([*wk.values(), ZERO])
    out['week_bars'] = [dict(label=names[k], value=wk.get(k, ZERO), pct=_pct(wk.get(k, ZERO), wpeak))
                        for k in (2, 3, 4, 5, 6, 7, 1)]

    # ---- Top customers (30 days) ----
    topc = list(Bill.objects.filter(status='FINAL', date__gte=today - timedelta(days=29))
                .exclude(customer_name__iexact='cash').exclude(customer_name='')
                .values('customer_name').annotate(s=Sum('items__total'), n=Count('id', distinct=True))
                .order_by('-s')[:5])
    cpeak = max([dec(t['s']) for t in topc] + [ZERO])
    out['top_customers'] = [dict(name=t['customer_name'], amount=dec(t['s']), bills=t['n'],
                                 pct=_pct(dec(t['s']), cpeak)) for t in topc]

    # ---- Reorder suggestion ----
    low = (Product.objects.exclude(is_active=False).filter(stock_qty__lte=F('reorder_level'))
           .select_related('unit').prefetch_related('unit_conversions__unit').order_by('stock_qty')[:6])
    out['reorder_rows'] = [dict(name=p.name or str(p), stock=p.stock_display(),
                                need=max(p.reorder_level * 2 - p.stock_qty, p.reorder_level),
                                unit=p.unit.name if p.unit_id else '') for p in low]

    # ---- Recent activity ----
    acts = []
    for b in final.prefetch_related('items').order_by('-date', '-id')[:5]:
        acts.append(dict(date=b.date, icon='🧾', text=f"Bill {b.display_number} · {b.customer_name} · ₹{b.grand_total}",
                         url=f"/bills/{b.pk}/invoice/"))
    for r in SalesReturn.objects.select_related('bill').order_by('-date', '-id')[:3]:
        acts.append(dict(date=r.date, icon='↩️', text=f"Return {r.note_number} · {r.bill.display_number}",
                         url=f"/sales_returns/{r.pk}/note/"))
    for r in CustomerPayment.objects.select_related('customer').order_by('-date', '-id')[:3]:
        acts.append(dict(date=r.date, icon='💰', text=f"Receipt ₹{r.amount} · {r.customer.name}",
                         url=f"/customers/{r.customer_id}/ledger/"))
    acts.sort(key=lambda a: a['date'] or today, reverse=True)
    out['activity'] = acts[:8]

    # ---- Today's Schedule H / H1 / X bills ----
    out['schedule_rows'] = [dict(inv=i.bill.display_number, product=i.product.name if i.product_id else '-',
                                 schedule=i.product.schedule if i.product_id else '',
                                 patient=i.bill.patient_name, doctor=i.bill.doctor_name)
                            for i in InvoiceItem.objects.filter(bill__status='FINAL', bill__date=today,
                                                                product__schedule__in=['H', 'H1', 'X'])
                            .select_related('bill', 'product').order_by('-id')[:8]]

    # ---- Staff/owner only: profit and stock value ----
    if include_money:
        out['money'] = True
        for key, start in (('today', today), ('month', month_start)):
            sales = dec(Sale.objects.filter(date__gte=start, date__lte=today).aggregate(s=Sum('taxable_value'))['s'])
            out[f'profit_{key}'] = sales - _cogs(start, today)
        out['stock_value'] = sum((dec(rate) * qty for qty, rate in
                                  Batch.objects.filter(qty__gt=0, purchase_item__isnull=False)
                                  .values_list('qty', 'purchase_item__rate')), ZERO)

    # ---- Supplier price jumps: latest price 10%+ above that supplier's previous price (last 60 days) ----
    recent = (PurchaseItem.objects.filter(purchase__invoice_date__gte=today - timedelta(days=60))
              .values_list('product_id', flat=True).distinct())
    out['price_jumps'] = supplier_prices.price_jumps(
        supplier_prices.rows_from_items(PurchaseItem.objects.filter(product_id__in=list(recent))), today)
    return out
